#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Разбор «График ППР АКМ на 2026г. РЭК.xlsx» (прислан заказчиком 25.09.2026).

Задача: полностью снять содержимое листа «РЭК» и выгрузить его в «длинный»
CSV так, чтобы ни одна строка исходника не потерялась.

Особенности исходника:
  * один лист «РЭК», данные в A9:H35 (шапка — строка 9, титул — строка 7);
  * объединённые ячейки: A10:A35 (подразделение), B.. (месяц), E/F/G/H (даты
    этапов) — значение верхней левой ячейки распространяется на весь блок;
  * колонка F — текст вида «до 900 13.01.2026», а не дата: разбираем в
    (пометка «до 900», дата);
  * даты приходят как datetime (эпоха Excel 1899-12-30 уже применена openpyxl).

Запуск (openpyxl ставится в .venv один раз):
    .venv/bin/pip install openpyxl
    .venv/bin/python analysis/to_ppr/parse_ppr.py
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "График ППР АКМ на 2026г. РЭК.xlsx"
OUT = ROOT / "analysis" / "to_ppr" / "ppr_schedule.csv"

SHEET = "РЭК"
HEADER_ROW = 9
COLS = {  # колонка исходника -> смысл
    "A": "подразделение",
    "B": "месяц_ппр",
    "C": "объект",
    "D": "кол_датчиков_шт",
}
# этапы: колонка -> имя этапа
STAGES = [
    ("E", "начало_демонтажа_датчиков"),
    ("F", "предоставление_в_ОМ_на_ППР_и_поверку"),
    ("G", "вывоз_датчиков_из_ОМ"),
    ("H", "сдача_работ_комиссии"),
]
F_RE = re.compile(r"^(?P<note>.*?)\s*(?P<date>\d{2}\.\d{2}\.\d{4})$")


def merged_lookup(ws):
    """(row, col) -> (верхняя-левая ячейка диапазона, флаг «из merge»)."""
    look = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                look[(row, col)] = (
                    rng.min_row,
                    rng.min_col,
                    (row, col) != (rng.min_row, rng.min_col),
                )
    return look


def cell_value(ws, look, row, col_letter):
    """Значение ячейки с учётом объединений: (значение, ячейка-источник, из_merge)."""
    col = openpyxl.utils.column_index_from_string(col_letter)
    if (row, col) in look:
        src_row, src_col, from_merge = look[(row, col)]
    else:
        src_row, src_col, from_merge = row, col, False
    cell = ws.cell(row=src_row, column=src_col)
    return cell.value, cell.coordinate, from_merge


def render_raw(value):
    """Как значение выглядит в файле (даты — ДД.ММ.ГГГГ)."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def to_iso(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return ""


def parse_f(raw: str):
    """«до 900 13.01.2026» -> («2026-01-13», «до 900»)."""
    raw = raw.strip()
    m = F_RE.match(raw)
    if not m:
        return "", raw
    note = m.group("note").strip()
    d, mo, y = m.group("date").split(".")
    return f"{y}-{mo}-{d}", note


def main() -> int:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[SHEET]
    look = merged_lookup(ws)

    # строки данных: те, где в колонке C («Наименование коллектора») есть значение
    rows = [
        r
        for r in range(HEADER_ROW + 1, ws.max_row + 1)
        if cell_value(ws, look, r, "C")[0] not in (None, "")
    ]

    # пакет работ = группа строк, разделяющих одну ячейку E (начало демонтажа)
    batch_of_row: dict[int, str] = {}
    seen: dict[tuple[int, int], str] = {}
    for r in rows:
        col_e = openpyxl.utils.column_index_from_string("E")
        key = look.get((r, col_e), (r, col_e, False))[:2]
        if key not in seen:
            seen[key] = f"P{len(seen) + 1:02d}"
        batch_of_row[r] = seen[key]

    out_rows = []
    for r in rows:
        base = {}
        for col_letter, name in COLS.items():
            value, coord, from_merge = cell_value(ws, look, r, col_letter)
            base[name] = render_raw(value)
            base[f"{name}_ячейка"] = coord
            base[f"{name}_из_merge"] = "да" if from_merge else "нет"
        for col_letter, stage in STAGES:
            value, coord, from_merge = cell_value(ws, look, r, col_letter)
            raw = render_raw(value)
            if col_letter == "F":
                iso, note = parse_f(raw)
            else:
                iso, note = to_iso(value), ""
            out_rows.append(
                {
                    "строка_листа": r,
                    "пакет_id": batch_of_row[r],
                    "подразделение": base["подразделение"],
                    "месяц_ппр": base["месяц_ппр"],
                    "объект": base["объект"],
                    "кол_датчиков_шт": base["кол_датчиков_шт"],
                    "этап": stage,
                    "дата_этапа": iso,
                    "пометка_из_файла": note,
                    "значение_как_в_файле": raw,
                    "ячейка": coord,
                    "из_объединённой_ячейки": "да" if from_merge else "нет",
                    "месяц_из_merge": base["месяц_ппр_из_merge"],
                }
            )

    fields = [
        "строка_листа",
        "пакет_id",
        "подразделение",
        "месяц_ппр",
        "объект",
        "кол_датчиков_шт",
        "этап",
        "дата_этапа",
        "пометка_из_файла",
        "значение_как_в_файле",
        "ячейка",
        "из_объединённой_ячейки",
        "месяц_из_merge",
    ]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, quoting=csv.QUOTE_ALL)
        w.writeheader()
        w.writerows(out_rows)

    # ---- сводка для инвентаризации ----------------------------------------
    objects = {r: cell_value(ws, look, r, "C")[0] for r in rows}
    qty = {r: cell_value(ws, look, r, "D")[0] for r in rows}
    print(f"строк с данными (объектов): {len(rows)}")
    print(f"пакетов работ: {len(set(batch_of_row.values()))}")
    print(f"строк в выгрузке: {len(out_rows)} (= объекты x {len(STAGES)} этапа)")
    print(f"сумма «Кол-во, шт.»: {sum(qty.values())}")
    dates = [x["дата_этапа"] for x in out_rows if x["дата_этапа"]]
    print(f"диапазон дат: {min(dates)} .. {max(dates)}")
    months = {}
    for r in rows:
        m = cell_value(ws, look, r, "B")[0]
        months.setdefault(m, []).append(qty[r])
    for m, qs in months.items():
        print(f"  {m}: объектов {len(qs)}, датчиков {sum(qs)}")
    empty = [x for x in out_rows if not x["дата_этапа"]]
    print(f"этапов без даты: {len(empty)}")
    for x in empty:
        print("  ", x["строка_листа"], x["объект"], x["этап"], repr(x["значение_как_в_файле"]))
    print("объекты:", ", ".join(str(objects[r]) for r in rows))
    return 0


if __name__ == "__main__":
    sys.exit(main())
