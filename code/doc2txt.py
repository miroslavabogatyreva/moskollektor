#!/usr/bin/env python
"""Извлекает текст из офисных файлов. Usage: doc2txt.py <file> [maxchars]"""
import sys, subprocess, zipfile, re, os

def docx(p):
    import docx as d
    doc = d.Document(p)
    out = [x.text for x in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            out.append(" | ".join(c.text.replace("\n", " ") for c in r.cells))
    return "\n".join(out)

def pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(p).slides, 1):
        out.append(f"\n--- слайд {i} ---")
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if sh.has_table:
                for r in sh.table.rows:
                    out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)

def xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n=== лист: {ws.title} ({ws.max_row}x{ws.max_column}) ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 400: out.append("... (обрезано)"); break
            cells = [str(c) for c in row if c is not None]
            if cells: out.append(" | ".join(cells))
    return "\n".join(out)

def pdf(p):
    return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout

def legacy(p):  # .doc .rtf .xls
    r = subprocess.run(["textutil", "-convert", "txt", "-stdout", p], capture_output=True, text=True)
    if r.stdout.strip(): return r.stdout
    return raw_zip(p)

def raw_zip(p):  # vsdx и всё, что не распарсилось
    try:
        z = zipfile.ZipFile(p)
    except Exception as e:
        return f"[не читается: {e}]"
    out = []
    for n in z.namelist():
        if n.endswith(".xml"):
            t = re.sub(r"<[^>]+>", " ", z.read(n).decode("utf8", "ignore"))
            t = re.sub(r"\s+", " ", t).strip()
            if len(t) > 40: out.append(t)
    return "\n".join(out)

HANDLERS = {".docx": docx, ".dotx": docx, ".pptx": pptx, ".ppt": raw_zip,
            ".xlsx": xlsx, ".xlsm": xlsx, ".pdf": pdf,
            ".doc": legacy, ".rtf": legacy, ".xls": legacy,
            ".vsdx": raw_zip, ".vsd": raw_zip, ".txt": lambda p: open(p, encoding="utf8", errors="ignore").read()}

def extract(p):
    ext = os.path.splitext(p)[1].lower()
    fn = HANDLERS.get(ext)
    if not fn: return f"[нет обработчика для {ext}]"
    try: return fn(p)
    except Exception as e: return f"[ошибка {type(e).__name__}: {e}]"

if __name__ == "__main__":
    lim = int(sys.argv[2]) if len(sys.argv) > 2 else 120000
    txt = re.sub(r"\n{3,}", "\n\n", extract(sys.argv[1]))
    print(txt[:lim])
    if len(txt) > lim: print(f"\n[...обрезано, всего {len(txt)} символов]")
